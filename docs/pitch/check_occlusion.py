#!/usr/bin/env python3
"""Flags any text that is hidden, covered or hard to read in an exported deck PDF.

    python3 docs/pitch/check_occlusion.py docs/pitch/CargoFlow-pitch.pdf
    python3 docs/pitch/check_occlusion.py CargoFlow-pitch-from-pptx.pdf      # a LibreOffice render of the PPTX

For every text span on every page (PyMuPDF text trace, in paint order) it reports:
  * IMAGE    text that overlaps a raster image (text must never sit on, or under, a picture);
  * COVERED  a filled shape painted AFTER the text that overlaps it (a box, tile, bar or stamp hiding text);
  * CONTRAST the fill directly underneath the text gives a contrast ratio below 3:1 (e.g. a dark box behind dark text).
Thin fills (<= 5 px tall, such as link underlines) are ignored. Exit code 1 if anything is found.
Needs PyMuPDF:  pip install pymupdf
"""
import sys
import pymupdf as fitz


def lum(rgb):
    def ch(v):
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(c) for c in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    x, y = sorted((lum(a), lum(b)), reverse=True)
    return (x + 0.05) / (y + 0.05)


def rgb_of(color):
    if isinstance(color, int):
        return ((color >> 16 & 255) / 255, (color >> 8 & 255) / 255, (color & 255) / 255)
    if color is None:
        return (0, 0, 0)
    if len(color) == 1:
        return (color[0],) * 3
    if len(color) == 4:  # cmyk
        c, m, y, k = color
        return ((1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k))
    return tuple(color[:3])


def area(r):
    return max(0.0, r.width) * max(0.0, r.height)


def check(path):
    doc = fitz.open(path)
    findings = []
    spans_total = 0
    for pno, page in enumerate(doc, 1):
        log = page.get_bboxlog()
        images = []
        for i, (kind, r) in enumerate(log):
            if kind != "fill-image":
                continue
            ir = fitz.Rect(r)
            # the bbox log ignores clipping: an object-fit:cover image is clipped by its frame, which is the
            # fill or stroke painted just before it; intersect with that frame when there is one
            for kind2, r2 in reversed(log[max(0, i - 3):i]):
                fr = fitz.Rect(r2)
                if kind2 in ("fill-path", "stroke-path") and area(fr & ir) >= 0.5 * area(fr) and area(fr) < 0.98 * area(ir):
                    ir = ir & fr
                    break
            images.append((i, ir))
        fills = []
        for d in page.get_drawings(extended=False):
            if d.get("fill") is None or d.get("type") not in ("f", "fs"):
                continue
            r = fitz.Rect(d["rect"])
            if r.height <= 5 or r.width <= 5:
                continue
            fills.append((d["seqno"], r, rgb_of(d["fill"]), d.get("fill_opacity", 1) or 1))
        page_bg = (1, 1, 1)
        for span in page.get_texttrace():
            text = "".join(chr(c[0]) for c in span["chars"]).strip()
            if not text or span.get("opacity", 1) == 0 or span.get("type", 0) == 3:
                continue
            spans_total += 1
            sb = fitz.Rect(span["bbox"])
            if area(sb) < 1:
                continue
            seq = span["seqno"]
            color = rgb_of(span["color"])
            label = f"p{pno} '{text[:40]}'"
            for _, ir in images:
                ov = area(sb & ir)
                if ov > 0.02 * area(sb):
                    findings.append(f"IMAGE    {label} overlaps an image at {tuple(round(v) for v in ir)}")
                    break
            under = page_bg
            under_seq = -1
            for fseq, fr, fcol, op in fills:
                ov = area(sb & fr)
                if ov <= 0:
                    continue
                if fseq > seq and ov > 0.10 * area(sb) and op > 0.2:
                    findings.append(f"COVERED  {label} is covered by a filled shape {tuple(round(v) for v in fr)} painted after it")
                elif fseq < seq and fr.contains(fitz.Point((sb.x0 + sb.x1) / 2, (sb.y0 + sb.y1) / 2)) and fseq > under_seq and op > 0.5:
                    under, under_seq = fcol, fseq
            cr = contrast(color, under)
            if cr < 3:
                findings.append(f"CONTRAST {label} text {tuple(round(c * 255) for c in color)} on {tuple(round(c * 255) for c in under)}: {cr:.2f}:1")
    return doc.page_count, spans_total, findings


def check_pptx(path):
    """Same idea on the PowerPoint shapes (python-pptx): no text frame may overlap a picture, no filled
    shape may sit above a text frame, and two text frames may not overlap."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    findings, n = [], 0
    for sno, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)  # z-order, bottom first
        boxes = []
        for z, sh in enumerate(shapes):
            if sh.width is None:
                continue
            # shrink by 3 px (1 px = 6350 EMU at 1920 px wide) so boxes that merely touch do not count
            t = 3 * 6350
            r = fitz.Rect(sh.left + t, sh.top + t, sh.left + sh.width - t, sh.top + sh.height - t)
            has_text = sh.has_text_frame and sh.text_frame.text.strip() != ""
            is_pic = sh.shape_type == MSO_SHAPE_TYPE.PICTURE
            filled = False
            if not has_text and not is_pic and sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    filled = sh.fill.type is not None and sh.fill.type == 1  # solid
                except Exception:
                    filled = False
            boxes.append((z, sh, r, has_text, is_pic, filled))
        for z, sh, r, has_text, _, _ in boxes:
            if not has_text:
                continue
            n += 1
            label = f"slide {sno} '{sh.text_frame.text.strip()[:40]}'"
            for z2, sh2, r2, t2, pic2, fill2 in boxes:
                if z2 == z or area(r & r2) <= 0:
                    continue
                if pic2:
                    findings.append(f"IMAGE    {label} overlaps picture '{sh2.name}'")
                elif t2 and z2 > z:
                    findings.append(f"OVERLAP  {label} overlaps text '{sh2.text_frame.text.strip()[:30]}'")
                elif fill2 and z2 > z and r2.height > 5 * 12700 and r2.width > 5 * 12700:
                    findings.append(f"COVERED  {label} is under filled shape '{sh2.name}'")
    return len(prs.slides), n, findings


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "docs/pitch/CargoFlow-pitch.pdf"
    pages, spans, findings = (check_pptx if target.endswith(".pptx") else check)(target)
    for f in findings:
        print(f)
    unit = "text frames" if target.endswith(".pptx") else "text spans"
    print(f"{target}: {pages} pages, {spans} {unit} checked, {len(findings)} finding(s)")
    sys.exit(1 if findings else 0)
